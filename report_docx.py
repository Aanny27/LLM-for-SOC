"""Sinh báo cáo Word (.docx) từ bảng analysis trong soc_results.db.

Dùng: data = build_report(DB_PATH, from_ts, to_ts)  ->  bytes (hoặc None nếu không có dữ liệu)
"""
import io
import json
import sqlite3
from collections import Counter
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

SEV_ORDER = ["CRITICAL", "HIGH", "MEDIUM", "LOW"]
SEV_COLOR = {"CRITICAL": "C0392B", "HIGH": "E67E22", "MEDIUM": "D4AC0D", "LOW": "7F8C8D"}
VERDICT_VI = {
    "true_positive": "Dương tính thật",
    "false_positive": "Dương tính giả",
}
MAX_DETAIL = 15   # số sự cố nghiêm trọng được mô tả chi tiết
MAX_TABLE = 30    # số dòng tối đa của bảng tổng hợp


def _fetch(db_path, from_ts, to_ts, hide_sca):
    # Cùng bộ lọc với /api/v1/soc/alerts
    cond = ["rule_description NOT LIKE 'CVE-%affects%'"]
    params = []
    if hide_sca:
        cond.append("rule_description NOT LIKE '%CIS Microsoft%' "
                    "AND rule_description NOT LIKE '%SCA summary%'")
    if from_ts:
        cond.append("timestamp >= ?")
        params.append(from_ts)
    if to_ts:
        cond.append("timestamp <= ?")
        params.append(to_ts)
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    try:
        rows = conn.execute(
            f"SELECT * FROM analysis WHERE {' AND '.join(cond)} ORDER BY id DESC", params
        ).fetchall()
    finally:
        conn.close()
    return [dict(r) for r in rows]


def _shade(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _table(doc, headers, data, sev_col=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(c, "1F2937")
    for row in data:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v if v not in (None, "") else "—")
            if sev_col is not None and i == sev_col and str(v) in SEV_COLOR:
                _shade(cells[i], SEV_COLOR[str(v)])
                for r in cells[i].paragraphs[0].runs:
                    r.bold = True
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    return t


def _actions(raw):
    try:
        val = json.loads(raw) if raw else []
    except (TypeError, ValueError):
        return [str(raw)]
    if isinstance(val, list):
        return [str(x) for x in val]
    return [str(val)]


def build_report(db_path, from_ts=None, to_ts=None, hide_sca=True):
    rows = _fetch(db_path, from_ts, to_ts, hide_sca)
    if not rows:
        return None

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    # --- Tiêu đề ---
    title = doc.add_heading("BÁO CÁO PHÂN TÍCH CẢNH BÁO AN NINH (SOC)", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rng = f"{from_ts or 'đầu dữ liệu'} → {to_ts or 'hiện tại'}"
    doc.add_paragraph(f"Thời gian xuất báo cáo: {datetime.now():%H:%M:%S %d/%m/%Y}")
    doc.add_paragraph(f"Khoảng thời gian: {rng}")
    doc.add_paragraph("Nguồn: hệ thống SOC tích hợp AI (Wazuh + mô hình cục bộ)")

    # --- 1. Tổng quan ---
    sev_cnt = Counter(r["severity"] for r in rows)
    agents = {r["agent_name"] for r in rows}
    labeled = [r for r in rows if r.get("verdict") in VERDICT_VI]
    tp = sum(1 for r in labeled if r["verdict"] == "true_positive")
    tp_rate = f"{tp / len(labeled) * 100:.0f}%" if labeled else "chưa có nhãn"

    doc.add_heading("1. Tổng quan", level=1)
    _table(doc, ["Chỉ số", "Giá trị"], [
        ("Tổng số sự cố", len(rows)),
        ("Mức nghiêm trọng (CRITICAL)", sev_cnt.get("CRITICAL", 0)),
        ("Mức cao (HIGH)", sev_cnt.get("HIGH", 0)),
        ("Số máy bị ảnh hưởng", len(agents)),
        ("Tỷ lệ dương tính thật (theo nhãn người)", tp_rate),
    ])

    # --- 2. Phân bố theo mức ---
    doc.add_heading("2. Phân bố theo mức nghiêm trọng", level=1)
    _table(doc, ["Mức", "Số lượng", "Tỷ lệ"],
           [(s, sev_cnt.get(s, 0), f"{sev_cnt.get(s, 0) / len(rows) * 100:.1f}%")
            for s in SEV_ORDER], sev_col=0)

    # --- 3. Top quy tắc & máy ---
    doc.add_heading("3. Quy tắc và máy xuất hiện nhiều nhất", level=1)
    doc.add_paragraph("Quy tắc:")
    top_rules = Counter((r["rule_id"], r["rule_description"]) for r in rows).most_common(5)
    _table(doc, ["Mã quy tắc", "Mô tả", "Số lần"],
           [(rid, desc, n) for (rid, desc), n in top_rules])
    doc.add_paragraph("")
    doc.add_paragraph("Máy:")
    _table(doc, ["Máy", "Số sự cố"], Counter(r["agent_name"] for r in rows).most_common(5))

    # --- 4. Bảng sự cố theo rủi ro ---
    doc.add_heading(f"4. Sự cố có điểm rủi ro cao nhất (tối đa {MAX_TABLE})", level=1)
    ranked = sorted(rows, key=lambda r: (r.get("risk_score") or 0), reverse=True)[:MAX_TABLE]
    _table(doc, ["Thời gian", "Máy", "Mô tả", "Mức", "Rủi ro", "Nhận định AI"],
           [(r["timestamp"], r["agent_name"], r["rule_description"], r["severity"],
             r.get("risk_score"), VERDICT_VI.get(r.get("ai_verdict"), "Chưa kết luận"))
            for r in ranked], sev_col=3)

    # --- 5. Chi tiết sự cố nghiêm trọng ---
    severe = [r for r in ranked if r["severity"] in ("CRITICAL", "HIGH")][:MAX_DETAIL]
    doc.add_heading("5. Chi tiết các sự cố nghiêm trọng", level=1)
    if not severe:
        doc.add_paragraph("Không có sự cố mức CRITICAL/HIGH trong khoảng thời gian này.")
    for r in severe:
        doc.add_heading(f"#{r['id']} — {r['rule_description']}", level=2)
        meta = (f"Thời gian: {r['timestamp']} | Máy: {r['agent_name']} | "
                f"Mức: {r['severity']} | Rủi ro: {r.get('risk_score') or '—'} | "
                f"MITRE: {r.get('mitre_technique') or '—'}")
        if (r.get("alert_count") or 1) > 1:
            meta += f" | Gộp {r['alert_count']} cảnh báo (IP nguồn: {r.get('source_ip') or '—'})"
        doc.add_paragraph(meta)
        doc.add_paragraph(f"Nhận định AI: {VERDICT_VI.get(r.get('ai_verdict'), 'Chưa kết luận')}"
                          + (f" — {r['ai_verdict_reason']}" if r.get("ai_verdict_reason") else ""))
        if r.get("analysis"):
            doc.add_paragraph(r["analysis"])
        acts = _actions(r.get("recommended_actions"))
        if acts:
            doc.add_paragraph("Khuyến nghị xử lý:")
            for a in acts:
                doc.add_paragraph(a, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
